"""Local document ingestion and evidence-location search for the V13 knowledge base.

Imported text remains source text. Parsing never creates accepted parameter claims
or silently binds a store SKU to a manufacturer model.
"""
from __future__ import annotations
import hashlib, json, mimetypes, os, re, sqlite3, tempfile, uuid, zipfile
from datetime import date, datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any
from .catalog_v12 import _validate_catalog
from .database_v9 import DEFAULT_DB, open_database
from .knowledge_schema_v13 import open_knowledge_readonly, validate_knowledge
from .knowledge_v13 import register_source

INGEST_SCHEMA_VERSION="1"
PARSER_VERSION="1.0"
MAX_SOURCE_BYTES=100*1024*1024
MAX_PDF_PAGES=5000
MAX_EXCEL_ROWS=100_000
MAX_EXCEL_COLUMNS=512
MAX_ARCHIVE_ENTRIES=20_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES=250*1024*1024
MAX_SEGMENT_CHARS=30_000
SUPPORTED_SUFFIXES={".pdf",".docx",".md",".markdown",".txt",".xlsx",".xlsm",".xls"}
_DDL=(
"""CREATE TABLE knowledge_ingest_meta(key TEXT PRIMARY KEY,value TEXT NOT NULL)""",
"""CREATE TABLE knowledge_document_segments(
id TEXT PRIMARY KEY,source_id TEXT NOT NULL REFERENCES knowledge_sources(id) ON DELETE RESTRICT,
parser_version TEXT NOT NULL,segment_key TEXT NOT NULL,
segment_type TEXT NOT NULL CHECK(segment_type IN ('pdf_page','pdf_line','docx_paragraph','docx_table_row',
'docx_section','markdown_line','markdown_section','text_line','text_block','excel_row')),
ordinal INTEGER NOT NULL CHECK(ordinal>=0),pdf_page INTEGER CHECK(pdf_page IS NULL OR pdf_page>0),
printed_page TEXT NOT NULL DEFAULT '',sheet_name TEXT NOT NULL DEFAULT '',cell_range TEXT NOT NULL DEFAULT '',
line_start INTEGER CHECK(line_start IS NULL OR line_start>0),
line_end INTEGER CHECK(line_end IS NULL OR line_end>=line_start),heading_path TEXT NOT NULL DEFAULT '',
text TEXT NOT NULL,text_sha256 TEXT NOT NULL CHECK(length(text_sha256)=64 AND text_sha256 NOT GLOB '*[^0-9a-f]*'),
region_json TEXT NOT NULL DEFAULT '{}',created_utc TEXT NOT NULL,
UNIQUE(source_id,parser_version,segment_key))""",
"""CREATE TABLE knowledge_ingest_runs(
event_id INTEGER PRIMARY KEY AUTOINCREMENT,run_id TEXT NOT NULL UNIQUE,
source_id TEXT NOT NULL REFERENCES knowledge_sources(id) ON DELETE RESTRICT,
parser_version TEXT NOT NULL,status TEXT NOT NULL CHECK(status IN ('indexed','failed')),
segment_count INTEGER NOT NULL CHECK(segment_count>=0),storage_relpath TEXT NOT NULL,
metadata_json TEXT NOT NULL DEFAULT '{}',warning_json TEXT NOT NULL DEFAULT '[]',
error_summary TEXT NOT NULL DEFAULT '',created_utc TEXT NOT NULL)""",
"CREATE INDEX idx_doc_segments_source_page ON knowledge_document_segments(source_id,pdf_page,ordinal)",
"CREATE INDEX idx_doc_segments_source_sheet ON knowledge_document_segments(source_id,sheet_name,cell_range)",
"CREATE INDEX idx_doc_segments_type ON knowledge_document_segments(segment_type,ordinal)",
"CREATE INDEX idx_ingest_runs_source ON knowledge_ingest_runs(source_id,event_id DESC)",
"CREATE INDEX idx_ingest_runs_status ON knowledge_ingest_runs(status,event_id DESC)",
"CREATE TRIGGER knowledge_ingest_meta_no_update BEFORE UPDATE ON knowledge_ingest_meta BEGIN SELECT RAISE(ABORT,'ingestion metadata is append-only'); END",
"CREATE TRIGGER knowledge_ingest_meta_no_delete BEFORE DELETE ON knowledge_ingest_meta BEGIN SELECT RAISE(ABORT,'ingestion metadata is append-only'); END",
"CREATE TRIGGER knowledge_document_segments_no_update BEFORE UPDATE ON knowledge_document_segments BEGIN SELECT RAISE(ABORT,'document segments are append-only'); END",
"CREATE TRIGGER knowledge_document_segments_no_delete BEFORE DELETE ON knowledge_document_segments BEGIN SELECT RAISE(ABORT,'document segments are append-only'); END",
"CREATE TRIGGER knowledge_ingest_runs_no_update BEFORE UPDATE ON knowledge_ingest_runs BEGIN SELECT RAISE(ABORT,'ingestion runs are append-only'); END",
"CREATE TRIGGER knowledge_ingest_runs_no_delete BEFORE DELETE ON knowledge_ingest_runs BEGIN SELECT RAISE(ABORT,'ingestion runs are append-only'); END")
_REQUIRED_TABLES={"knowledge_ingest_meta","knowledge_document_segments","knowledge_ingest_runs"}
_REQUIRED_TRIGGERS={"knowledge_ingest_meta_no_update","knowledge_ingest_meta_no_delete",
"knowledge_document_segments_no_update","knowledge_document_segments_no_delete",
"knowledge_ingest_runs_no_update","knowledge_ingest_runs_no_delete"}
_PRECISE_TYPES={"pdf_line","docx_paragraph","docx_table_row","markdown_line","text_line","excel_row"}


def _now()->str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _id()->str:
    return uuid.uuid4().hex


def _json(value:Any)->str:
    return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(",",":"),allow_nan=False)


def _safe_path(value:str|Path,label:str)->Path:
    path=Path(value).expanduser().absolute()
    if path.is_symlink() or any(parent.is_symlink() for parent in path.parents):
        raise ValueError(f"{label}路径不能包含符号链接")
    return path


def _db_path(path:str|Path)->Path:
    return Path(path).expanduser().absolute()


def _storage_root(path:str|Path)->Path:
    root=_db_path(path).parent/"knowledge_sources"
    _safe_path(root,"资料存储")
    root.mkdir(parents=True,exist_ok=True,mode=0o700)
    if os.name=="posix": os.chmod(root,0o700)
    return root


def validate_ingestion(conn:sqlite3.Connection)->None:
    tables={r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    if not _REQUIRED_TABLES.issubset(tables):
        raise RuntimeError("文档检索表尚未初始化；请执行 knowledge_ingest_v13.py init")
    version=conn.execute("SELECT value FROM knowledge_ingest_meta WHERE key='schema_version'").fetchone()
    if not version or version[0]!=INGEST_SCHEMA_VERSION:
        raise RuntimeError("文档检索结构版本缺失或不受支持")
    triggers={r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='trigger'")}
    if not _REQUIRED_TRIGGERS.issubset(triggers):
        raise RuntimeError("文档检索保护触发器缺失；拒绝继续操作")


def initialize_ingestion(path:str|Path=DEFAULT_DB)->None:
    """Create a separately versioned, additive document index on top of V13."""
    with open_database(path) as conn:
        _validate_catalog(conn); validate_knowledge(conn); conn.execute("BEGIN IMMEDIATE")
        try:
            has_meta=conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='knowledge_ingest_meta'").fetchone()
            objects={r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type IN ('table','trigger')")}
            partial={n for n in objects if n in _REQUIRED_TABLES or n.startswith("knowledge_document_segments")
                     or n.startswith("knowledge_ingest_runs") or n.startswith("knowledge_ingest_meta")}
            if has_meta:
                version=conn.execute("SELECT value FROM knowledge_ingest_meta WHERE key='schema_version'").fetchone()
                if not version or version[0]!=INGEST_SCHEMA_VERSION: raise RuntimeError("检测到未知文档检索版本，拒绝覆盖")
                validate_ingestion(conn); conn.execute("COMMIT"); return
            if partial: raise RuntimeError("发现未标记版本或不完整的文档检索对象，拒绝自动接管")
            for statement in _DDL: conn.execute(statement)
            conn.execute("INSERT INTO knowledge_ingest_meta(key,value) VALUES('schema_version',?)",(INGEST_SCHEMA_VERSION,))
            validate_ingestion(conn); conn.execute("COMMIT")
        except BaseException:
            if conn.in_transaction: conn.execute("ROLLBACK")
            raise


def _sha256_file(path:Path)->tuple[str,int]:
    digest=hashlib.sha256(); size=0
    with path.open("rb") as stream:
        for chunk in iter(lambda:stream.read(1024*1024),b""):
            digest.update(chunk); size+=len(chunk)
    return digest.hexdigest(),size


def _stage_file(source:Path,root:Path,max_bytes:int)->tuple[Path,str,int,Path,str]:
    suffix=source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        if suffix==".doc": raise ValueError("不支持旧版 .doc；请在 Word 中另存为 .docx 后导入")
        raise ValueError(f"不支持的文件格式：{suffix or '无扩展名'}；支持 PDF、DOCX、MD、TXT、XLSX、XLSM、XLS")
    if not source.exists() or not source.is_file(): raise ValueError("原始资料文件不存在或不是普通文件")
    if source.stat().st_size<=0: raise ValueError("原始资料文件为空")
    if source.stat().st_size>max_bytes: raise ValueError(f"原始资料超过大小限制（{max_bytes} 字节）")
    staging=root/".staging"; _safe_path(staging,"临时资料")
    staging.mkdir(parents=True,exist_ok=True,mode=0o700)
    fd,name=tempfile.mkstemp(prefix="ingest-",suffix=suffix,dir=staging); tmp=Path(name)
    digest=hashlib.sha256(); size=0
    try:
        with source.open("rb") as src,os.fdopen(fd,"wb") as dst:
            for chunk in iter(lambda:src.read(1024*1024),b""):
                size+=len(chunk)
                if size>max_bytes: raise ValueError(f"原始资料超过大小限制（{max_bytes} 字节）")
                digest.update(chunk); dst.write(chunk)
            dst.flush(); os.fsync(dst.fileno())
        if size<=0: raise ValueError("原始资料文件为空")
        sha=digest.hexdigest(); final_dir=root/sha[:2]; _safe_path(final_dir,"资料分片目录")
        final_dir.mkdir(parents=True,exist_ok=True,mode=0o700)
        final=final_dir/f"{sha}{suffix}"
        relative=str(PurePosixPath("knowledge_sources")/sha[:2]/f"{sha}{suffix}")
        return tmp,sha,size,final,relative
    except BaseException:
        try: tmp.unlink(missing_ok=True)
        except OSError: pass
        raise


def _finalize_staged_file(tmp:Path,final:Path,sha:str,size:int)->None:
    _safe_path(final,"原始资料目标")
    if final.exists():
        if not final.is_file(): raise ValueError("原始资料目标不是普通文件")
        actual,actual_size=_sha256_file(final)
        if actual!=sha or actual_size!=size: raise RuntimeError("资料存储中已有指纹不一致的文件，拒绝覆盖")
        tmp.unlink(missing_ok=True); return
    os.replace(tmp,final)
    if os.name=="posix": os.chmod(final,0o600)


def _safe_archive(path:Path)->None:
    """Bound Office ZIP expansion before handing files to document libraries."""
    try:
        with zipfile.ZipFile(path) as archive:
            members=archive.infolist()
            if len(members)>MAX_ARCHIVE_ENTRIES: raise ValueError("Office 文件包含过多压缩条目")
            total=0
            for item in members:
                if item.file_size<0 or item.compress_size<0: raise ValueError("Office 文件压缩目录无效")
                total+=item.file_size
                if total>MAX_ARCHIVE_UNCOMPRESSED_BYTES: raise ValueError("Office 文件解压后体积超过安全限制")
                if item.file_size>5*1024*1024 and item.file_size/max(item.compress_size,1)>1000:
                    raise ValueError("Office 文件压缩比异常，拒绝解析")
            if archive.testzip() is not None: raise ValueError("Office 文件包含损坏的压缩条目")
    except zipfile.BadZipFile as exc:
        raise ValueError("Office 文件不是有效的 ZIP 容器或文件已损坏") from exc


def _segment(key:str,kind:str,ordinal:int,text:str,*,pdf_page:int|None=None,printed_page:str="",
             sheet_name:str="",cell_range:str="",line_start:int|None=None,line_end:int|None=None,
             heading_path:str="",region:dict[str,Any]|None=None)->dict[str,Any]|None:
    if not isinstance(text,str) or not text.strip(): return None
    clean=text.strip()
    return {"segment_key":key,"segment_type":kind,"ordinal":ordinal,"pdf_page":pdf_page,
            "printed_page":printed_page,"sheet_name":sheet_name,"cell_range":cell_range,
            "line_start":line_start,"line_end":line_end,"heading_path":heading_path,"text":clean,
            "text_sha256":hashlib.sha256(clean.encode("utf-8")).hexdigest(),"region_json":_json(region or {})}


def _bbox_region(bbox:Any,rect:Any,**extra:Any)->dict[str,Any]:
    width=max(float(rect.width),1.0); height=max(float(rect.height),1.0)
    values=[float(x) for x in (bbox or [rect.x0,rect.y0,rect.x1,rect.y1])]
    x0,y0,x1,y1=values
    return {"bbox":[round(x,2) for x in values],
            "normalized_bbox":[round((x0-float(rect.x0))/width,6),round((y0-float(rect.y0))/height,6),
                               round((x1-float(rect.x0))/width,6),round((y1-float(rect.y0))/height,6)],
            "coordinate_system":"pdf_points_top_left","page_size":[round(width,2),round(height,2)],**extra}


def _parse_pdf(path:Path)->tuple[list[dict[str,Any]],int,list[str]]:
    try: import fitz
    except ImportError as exc: raise RuntimeError("解析 PDF 需要 PyMuPDF：pip install 'PyMuPDF>=1.24,<2'") from exc
    try: document=fitz.open(str(path))
    except Exception as exc: raise ValueError(f"无法打开 PDF：{exc}") from exc
    segments=[]; warnings=[]
    try:
        if not document.is_pdf: raise ValueError("文件内容不是有效 PDF")
        if document.needs_pass: raise ValueError("PDF 已加密；请先获得授权并导出可读取的版本")
        count=len(document)
        if count<1 or count>MAX_PDF_PAGES: raise ValueError(f"PDF 页数无效或超过限制（1–{MAX_PDF_PAGES}）")
        ordinal=0; empty=[]
        for page_index in range(count):
            page=document[page_index]; rect=page.rect; lines=[]
            data=page.get_text("dict",sort=True)
            for block_index,block in enumerate(data.get("blocks",[]),start=1):
                if block.get("type",0)!=0: continue
                for block_line_index,line in enumerate(block.get("lines",[]),start=1):
                    spans=line.get("spans",[]); line_text="".join(str(s.get("text","")) for s in spans).strip()
                    if not line_text: continue
                    bbox=line.get("bbox")
                    if not bbox and spans:
                        boxes=[s.get("bbox") for s in spans if s.get("bbox")]
                        if boxes: bbox=[min(b[0] for b in boxes),min(b[1] for b in boxes),max(b[2] for b in boxes),max(b[3] for b in boxes)]
                    if bbox: lines.append(([float(x) for x in bbox],line_text,block_index,block_line_index))
            page_text=page.get_text("text",sort=True).strip()
            if not page_text and lines: page_text="\n".join(x[1] for x in lines)
            page_no=page_index+1
            if not page_text: empty.append(page_no); continue
            union=[min(x[0][0] for x in lines),min(x[0][1] for x in lines),
                   max(x[0][2] for x in lines),max(x[0][3] for x in lines)] if lines else [rect.x0,rect.y0,rect.x1,rect.y1]
            ordinal+=1
            item=_segment(f"pdf:{page_no}:page","pdf_page",ordinal,page_text,pdf_page=page_no,
                line_start=1 if lines else None,line_end=len(lines) if lines else None,
                region=_bbox_region(union,rect,precision="page_text_union",line_count=len(lines)))
            if item: segments.append(item)
            for line_no,(bbox,line_text,block_no,block_line_no) in enumerate(lines,start=1):
                ordinal+=1
                item=_segment(f"pdf:{page_no}:line:{line_no}","pdf_line",ordinal,line_text,
                    pdf_page=page_no,line_start=line_no,line_end=line_no,
                    region=_bbox_region(bbox,rect,block_index=block_no,block_line_index=block_line_no,precision="line"))
                if item: segments.append(item)
        if empty:
            warnings.append(f"{len(empty)} 页没有可提取文字层（PDF页码：{','.join(map(str,empty[:100]))}"
                            f"{'…' if len(empty)>100 else ''}）；可能是扫描页，本次未执行 OCR")
        return segments,count,warnings
    finally: document.close()


def _split_chunks(lines:list[tuple[int,str]],max_chars:int=MAX_SEGMENT_CHARS)->list[tuple[int,int,str]]:
    chunks=[]; current=[]; start=end=0; length=0
    for line_no,text in lines:
        added=len(text)+(1 if current else 0)
        if current and length+added>max_chars:
            chunks.append((start,end,"\n".join(current))); current=[]; length=0
        if not current: start=line_no
        current.append(text); end=line_no; length+=len(text)+(1 if len(current)>1 else 0)
    if current: chunks.append((start,end,"\n".join(current)))
    return chunks


def _parse_markdown(path:Path,plain_text:bool=False)->tuple[list[dict[str,Any]],None,list[str]]:
    try: content=path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc: raise ValueError("Markdown/TXT 必须使用 UTF-8 编码") from exc
    lines=content.splitlines(); segments=[]; warnings=[]; ordinal=0
    if not plain_text:
        headings=[]; section_lines=[]; section_heading=""; section_no=0
        def flush()->None:
            nonlocal ordinal,section_no,section_lines
            if not section_lines: return
            section_no+=1
            for chunk_no,(first,last,chunk) in enumerate(_split_chunks(section_lines),start=1):
                ordinal+=1
                item=_segment(f"markdown:section:{section_no}:chunk:{chunk_no}","markdown_section",ordinal,chunk,
                    line_start=first,line_end=last,heading_path=section_heading,
                    region={"line_start":first,"line_end":last,"precision":"section"})
                if item: segments.append(item)
            section_lines=[]
        for line_no,line in enumerate(lines,start=1):
            match=re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$",line)
            if match:
                flush(); level=len(match.group(1)); headings=headings[:level-1]; headings.append(match.group(2).strip())
                section_heading=" > ".join(headings); section_lines=[(line_no,line)]
            elif line.strip():
                if not section_lines: section_heading=" > ".join(headings)
                section_lines.append((line_no,line))
            if line.strip():
                ordinal+=1
                item=_segment(f"markdown:line:{line_no}","markdown_line",ordinal,line,line_start=line_no,line_end=line_no,
                    heading_path=" > ".join(headings),region={"line_start":line_no,"line_end":line_no,"precision":"line"})
                if item: segments.append(item)
        flush()
    else:
        blocks=[]; block_no=0
        def flush()->None:
            nonlocal ordinal,block_no,blocks
            if not blocks: return
            block_no+=1
            for chunk_no,(first,last,chunk) in enumerate(_split_chunks(blocks),start=1):
                ordinal+=1
                item=_segment(f"text:block:{block_no}:chunk:{chunk_no}","text_block",ordinal,chunk,
                    line_start=first,line_end=last,region={"line_start":first,"line_end":last,"precision":"paragraph_block"})
                if item: segments.append(item)
            blocks=[]
        for line_no,line in enumerate(lines,start=1):
            if not line.strip(): flush(); continue
            blocks.append((line_no,line)); ordinal+=1
            item=_segment(f"text:line:{line_no}","text_line",ordinal,line,line_start=line_no,line_end=line_no,
                region={"line_start":line_no,"line_end":line_no,"precision":"line"})
            if item: segments.append(item)
        flush()
    if not segments: warnings.append("文件中没有可索引的非空文本")
    return segments,None,warnings


def _parse_docx(path:Path)->tuple[list[dict[str,Any]],None,list[str]]:
    try:
        from docx import Document
        from docx.oxml.text.paragraph import CT_P
        from docx.oxml.table import CT_Tbl
        from docx.text.paragraph import Paragraph
        from docx.table import Table
    except ImportError as exc: raise RuntimeError("解析 DOCX 需要 python-docx：pip install 'python-docx>=1.1,<2'") from exc
    try: document=Document(str(path))
    except Exception as exc: raise ValueError(f"无法打开 DOCX：{exc}") from exc
    segments=[]; warnings=["DOCX 不保证固定分页；使用段落编号、表格行和标题路径定位"]
    headings=[]; section_parts=[]; section_heading=""; section_no=0; ordinal=0; para_no=0; table_no=0; body_no=0
    def flush()->None:
        nonlocal ordinal,section_no,section_parts
        if not section_parts: return
        section_no+=1
        for chunk_no,(first,last,chunk) in enumerate(_split_chunks(section_parts),start=1):
            ordinal+=1
            item=_segment(f"docx:section:{section_no}:chunk:{chunk_no}","docx_section",ordinal,chunk,
                line_start=first,line_end=last,heading_path=section_heading,
                region={"body_element_start":first,"body_element_end":last,"precision":"heading_section"})
            if item: segments.append(item)
        section_parts=[]
    for child in document.element.body.iterchildren():
        body_no+=1
        if isinstance(child,CT_P):
            para_no+=1; paragraph=Paragraph(child,document); text=paragraph.text.strip()
            if not text: continue
            style=paragraph.style.name if paragraph.style else ""
            match=re.match(r"^(?:Heading|标题)\s*(\d+)$",style,re.I)
            if match:
                flush(); level=max(1,min(int(match.group(1)),6)); headings=headings[:level-1]; headings.append(text)
                section_heading=" > ".join(headings)
            elif not section_parts: section_heading=" > ".join(headings)
            ordinal+=1
            item=_segment(f"docx:paragraph:{para_no}","docx_paragraph",ordinal,text,
                heading_path=" > ".join(headings),
                region={"paragraph_index":para_no,"body_element_index":body_no,"style":style,"precision":"paragraph"})
            if item:
                segments.append(item); section_parts.append((body_no,f"[段落 {para_no}] {text}"))
        elif isinstance(child,CT_Tbl):
            table_no+=1; table=Table(child,document); rows=list(table.rows); header=[]
            if not section_parts: section_heading=" > ".join(headings)
            for row_no,row in enumerate(rows,start=1):
                raw=[cell.text.replace("\n"," / ").strip() for cell in row.cells]
                if row_no==1 and sum(bool(x) for x in raw)>=2: header=raw
                values=[(col,value) for col,value in enumerate(raw,start=1) if value]
                if not values: continue
                parts=[]; details=[]
                for col,value in values:
                    label=header[col-1] if header and col<=len(header) and header[col-1] else f"C{col}"
                    parts.append(f"{label}={value}")
                    details.append({"cell":f"R{row_no}C{col}","header":label,"value":value})
                text=" | ".join(parts); ordinal+=1
                item=_segment(f"docx:table:{table_no}:row:{row_no}","docx_table_row",ordinal,text,
                    sheet_name=f"Table {table_no}",cell_range=f"R{row_no}C1:R{row_no}C{max(len(raw),1)}",
                    heading_path=" > ".join(headings),
                    region={"table_index":table_no,"row_index":row_no,"cells":details,"precision":"table_row"})
                if item:
                    segments.append(item); section_parts.append((body_no,f"[表格 {table_no} 行 {row_no}] {text}"))
    flush()
    if not segments: warnings.append("DOCX 中没有可索引的非空段落或表格")
    return segments,None,warnings


def _col_letter(number:int)->str:
    result=""
    while number:
        number,remainder=divmod(number-1,26); result=chr(65+remainder)+result
    return result


def _display(value:Any)->str:
    if value is None: return ""
    if isinstance(value,(datetime,date)): return value.isoformat()
    return str(value).strip()


def _parse_xlsx(path:Path)->tuple[list[dict[str,Any]],None,list[str]]:
    try: from openpyxl import load_workbook
    except ImportError as exc: raise RuntimeError("解析 XLSX/XLSM 需要 openpyxl：pip install 'openpyxl>=3.1,<4'") from exc
    try: workbook=load_workbook(str(path),read_only=True,data_only=False,keep_links=False)
    except Exception as exc: raise ValueError(f"无法打开 Excel 工作簿：{exc}") from exc
    segments=[]; warnings=[]; ordinal=0; formulas=0
    try:
        if len(workbook.worksheets)>500: raise ValueError("Excel 工作表数量超过安全限制")
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0)>MAX_EXCEL_ROWS or (sheet.max_column or 0)>MAX_EXCEL_COLUMNS:
                raise ValueError(f"工作表 {sheet.title} 行数或列数超过安全限制")
            header=[]
            for row_no,row in enumerate(sheet.iter_rows(),start=1):
                if row_no>MAX_EXCEL_ROWS: raise ValueError(f"工作表 {sheet.title} 行数超过安全限制")
                values=[]
                for cell in row:
                    value=_display(cell.value)
                    if cell.data_type=="f": formulas+=1
                    if value: values.append((cell.column,cell.coordinate,value))
                if not values: continue
                if not header and row_no<=20 and len(values)>=2:
                    header=[""]*max(x[0] for x in values)
                    for col,_,value in values: header[col-1]=value
                parts=[]; details=[]
                for col,coordinate,value in values:
                    label=header[col-1] if col<=len(header) and header[col-1] else coordinate
                    parts.append(f"{label}={value}" if label!=coordinate else f"{coordinate}={value}")
                    details.append({"cell":coordinate,"header":label,"value":value})
                cell_range=f"{_col_letter(min(x[0] for x in values))}{row_no}:{_col_letter(max(x[0] for x in values))}{row_no}"
                ordinal+=1
                item=_segment(f"xlsx:{sheet.title}:row:{row_no}","excel_row",ordinal," | ".join(parts),
                    sheet_name=sheet.title,cell_range=cell_range,line_start=row_no,line_end=row_no,
                    region={"sheet_name":sheet.title,"row_index":row_no,"cells":details,"precision":"worksheet_row"})
                if item: segments.append(item)
    finally: workbook.close()
    if formulas: warnings.append(f"发现 {formulas} 个公式单元格；只索引公式文本，不执行或重新计算公式")
    if not segments: warnings.append("Excel 工作簿中没有可索引的非空单元格")
    return segments,None,warnings


def _xls_value(cell:Any,datemode:int)->str:
    import xlrd
    if cell.ctype in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK): return ""
    if cell.ctype==xlrd.XL_CELL_DATE:
        try: return xlrd.xldate_as_datetime(cell.value,datemode).isoformat()
        except (ValueError,OverflowError): return str(cell.value)
    if cell.ctype==xlrd.XL_CELL_BOOLEAN: return "TRUE" if cell.value else "FALSE"
    if cell.ctype==xlrd.XL_CELL_ERROR: return f"EXCEL_ERROR:{cell.value}"
    return str(cell.value).strip()


def _parse_xls(path:Path)->tuple[list[dict[str,Any]],None,list[str]]:
    try: import xlrd
    except ImportError as exc: raise RuntimeError("解析旧版 XLS 需要 xlrd：pip install 'xlrd>=2.0.1,<3'") from exc
    try: workbook=xlrd.open_workbook(str(path),on_demand=True)
    except Exception as exc: raise ValueError(f"无法打开 XLS 工作簿：{exc}") from exc
    segments=[]; warnings=[]; ordinal=0
    try:
        if len(workbook.sheet_names())>500: raise ValueError("Excel 工作表数量超过安全限制")
        for sheet in workbook.sheets():
            if sheet.nrows>MAX_EXCEL_ROWS or sheet.ncols>MAX_EXCEL_COLUMNS:
                raise ValueError(f"工作表 {sheet.name} 行数或列数超过安全限制")
            header=[]
            for row_idx in range(sheet.nrows):
                row_no=row_idx+1; values=[]
                for col_idx in range(sheet.ncols):
                    value=_xls_value(sheet.cell(row_idx,col_idx),workbook.datemode)
                    if value: values.append((col_idx+1,value))
                if not values: continue
                if not header and row_no<=20 and len(values)>=2:
                    header=[""]*max(c for c,_ in values)
                    for col,value in values: header[col-1]=value
                parts=[]; details=[]
                for col,value in values:
                    coordinate=f"{_col_letter(col)}{row_no}"
                    label=header[col-1] if col<=len(header) and header[col-1] else coordinate
                    parts.append(f"{label}={value}" if label!=coordinate else f"{coordinate}={value}")
                    details.append({"cell":coordinate,"header":label,"value":value})
                cell_range=f"{_col_letter(min(c for c,_ in values))}{row_no}:{_col_letter(max(c for c,_ in values))}{row_no}"
                ordinal+=1
                item=_segment(f"xls:{sheet.name}:row:{row_no}","excel_row",ordinal," | ".join(parts),
                    sheet_name=sheet.name,cell_range=cell_range,line_start=row_no,line_end=row_no,
                    region={"sheet_name":sheet.name,"row_index":row_no,"cells":details,"precision":"worksheet_row"})
                if item: segments.append(item)
    finally: workbook.release_resources()
    if not segments: warnings.append("Excel 工作簿中没有可索引的非空单元格")
    return segments,None,warnings


def _parse_document(path:Path)->tuple[list[dict[str,Any]],int|None,list[str]]:
    suffix=path.suffix.lower()
    if suffix in {".docx",".xlsx",".xlsm"}: _safe_archive(path)
    if suffix==".pdf": return _parse_pdf(path)
    if suffix==".docx": return _parse_docx(path)
    if suffix in {".md",".markdown"}: return _parse_markdown(path)
    if suffix==".txt": return _parse_markdown(path,plain_text=True)
    if suffix in {".xlsx",".xlsm"}: return _parse_xlsx(path)
    if suffix==".xls": return _parse_xls(path)
    raise ValueError(f"不支持的文件格式：{suffix}")


def _register_source(path:str|Path,*,filename:str,source_kind:str,authority_level:str,sha256:str,size_bytes:int,
                     media_type:str,page_count:int|None,document_family:str,document_number:str,
                     revision_label:str,issue_date:str,storage_ref:str)->str:
    """Reuse matching provenance without mutating or upgrading existing source metadata."""
    with open_database(path) as conn:
        _validate_catalog(conn); validate_knowledge(conn)
        old=conn.execute("""SELECT id,original_filename,source_kind,authority_level,size_bytes,document_family,
            document_number,revision_label,issue_date,media_type,page_count FROM knowledge_sources WHERE sha256=?""",
            (sha256,)).fetchone()
    if old:
        expected=(filename,source_kind,authority_level,size_bytes,document_family,document_number,revision_label,issue_date)
        if tuple(old[1:9])!=expected:
            raise ValueError("同一 SHA-256 已存在但文件身份/来源元数据不一致；不会覆盖或升级原始来源记录")
        if old[9] and media_type and old[9]!=media_type: raise ValueError("同一 SHA-256 已存在但媒体类型不一致")
        if old[10] is not None and page_count is not None and old[10]!=page_count: raise ValueError("同一 SHA-256 已存在但页数不一致")
        return old[0]
    return register_source(path,filename=filename,source_kind=source_kind,authority_level=authority_level,
        sha256=sha256,size_bytes=size_bytes,media_type=media_type,page_count=page_count,
        document_family=document_family,document_number=document_number,revision_label=revision_label,
        issue_date=issue_date,storage_ref=storage_ref,metadata={})


def import_document(path:str|Path=DEFAULT_DB,*,file_path:str|Path,source_kind:str="other",
                    authority_level:str="user_supplied",document_family:str="",document_number:str="",
                    revision_label:str="",issue_date:str="",max_bytes:int=MAX_SOURCE_BYTES)->dict[str,Any]:
    """Copy, fingerprint, parse and index one local document; never create claims automatically."""
    if type(max_bytes) is not int or not 1<=max_bytes<=MAX_SOURCE_BYTES: raise ValueError(f"max_bytes 必须在 1 到 {MAX_SOURCE_BYTES} 之间")
    if source_kind not in {"manufacturer_document","supplier_document","purchase_record","physical_photo","lab_report","platform_export","internal_note","other"}:
        raise ValueError("资料类型无效")
    if authority_level not in {"unverified","user_supplied","manufacturer_controlled","independent_test","system_export"}:
        raise ValueError("来源级别无效")
    source=_safe_path(file_path,"原始资料")
    if not source.exists() or not source.is_file(): raise ValueError("原始资料文件不存在或不是普通文件")
    if source.name in {".",".."} or len(source.name)>255 or any(ord(c)<32 for c in source.name): raise ValueError("原始资料文件名无效")
    for value,label,limit in ((document_family,"文档系列",200),(document_number,"文件编号",200),(revision_label,"修订版",100),(issue_date,"发布日期",40)):
        if not isinstance(value,str) or len(value)>limit or any(ord(c)<32 for c in value): raise ValueError(f"{label}字段无效")
    if issue_date.strip():
        try: date.fromisoformat(issue_date.strip())
        except ValueError as exc: raise ValueError("发布日期必须使用 YYYY-MM-DD 格式") from exc
    initialize_ingestion(path)
    tmp,sha,size,final,storage_ref=_stage_file(source,_storage_root(path),max_bytes)
    try:
        segments,page_count,warnings=_parse_document(tmp)
        media_type=mimetypes.guess_type(source.name)[0] or "application/octet-stream"
        source_id=_register_source(path,filename=source.name,source_kind=source_kind,authority_level=authority_level,
            sha256=sha,size_bytes=size,media_type=media_type,page_count=page_count,
            document_family=document_family.strip(),document_number=document_number.strip(),
            revision_label=revision_label.strip(),issue_date=issue_date.strip(),storage_ref=storage_ref)
        _finalize_staged_file(tmp,final,sha,size)
        metadata={"file_format":source.suffix.lower().lstrip("."),"media_type":media_type,"page_count":page_count,
            "document_family":document_family.strip(),"document_number":document_number.strip(),
            "revision_label":revision_label.strip(),"issue_date":issue_date.strip(),"parser":"knowledge_ingest_v13"}
        with open_database(path) as conn:
            _validate_catalog(conn); validate_knowledge(conn); validate_ingestion(conn)
            old=conn.execute("""SELECT run_id,segment_count,storage_relpath,warning_json FROM knowledge_ingest_runs
                WHERE source_id=? AND parser_version=? AND status='indexed' ORDER BY event_id DESC LIMIT 1""",
                (source_id,PARSER_VERSION)).fetchone()
            if old:
                stored=_stored_path(path,old[2])
                actual_count=conn.execute("SELECT COUNT(*) FROM knowledge_document_segments WHERE source_id=? AND parser_version=?",
                    (source_id,PARSER_VERSION)).fetchone()[0]
                if stored.is_file():
                    stored_sha,stored_size=_sha256_file(stored)
                    if stored_sha!=sha or stored_size!=size: raise RuntimeError("已索引原始文件指纹不一致；拒绝静默覆盖")
                if stored.is_file() and actual_count==old[1]:
                    return {"run_id":old[0],"source_id":source_id,"filename":source.name,"sha256":sha,
                        "size_bytes":size,"parser_version":PARSER_VERSION,"segment_count":actual_count,
                        "storage_relpath":old[2],"warnings":json.loads(old[3]),"already_indexed":True}
            conn.execute("BEGIN IMMEDIATE")
            try:
                for item in segments:
                    conn.execute("""INSERT OR IGNORE INTO knowledge_document_segments(
                        id,source_id,parser_version,segment_key,segment_type,ordinal,pdf_page,printed_page,
                        sheet_name,cell_range,line_start,line_end,heading_path,text,text_sha256,region_json,created_utc)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (_id(),source_id,PARSER_VERSION,item["segment_key"],item["segment_type"],item["ordinal"],
                         item["pdf_page"],item["printed_page"],item["sheet_name"],item["cell_range"],
                         item["line_start"],item["line_end"],item["heading_path"],item["text"],
                         item["text_sha256"],item["region_json"],_now()))
                count=conn.execute("SELECT COUNT(*) FROM knowledge_document_segments WHERE source_id=? AND parser_version=?",
                    (source_id,PARSER_VERSION)).fetchone()[0]
                run_id=_id()
                conn.execute("""INSERT INTO knowledge_ingest_runs(run_id,source_id,parser_version,status,segment_count,
                    storage_relpath,metadata_json,warning_json,error_summary,created_utc) VALUES(?,?,?,?,?,?,?,?,?,?)""",
                    (run_id,source_id,PARSER_VERSION,"indexed",count,storage_ref,_json(metadata),_json(warnings),"",_now()))
                conn.execute("COMMIT")
            except BaseException:
                if conn.in_transaction: conn.execute("ROLLBACK")
                raise
        return {"run_id":run_id,"source_id":source_id,"filename":source.name,"sha256":sha,"size_bytes":size,
            "parser_version":PARSER_VERSION,"segment_count":count,"storage_relpath":storage_ref,
            "warnings":warnings,"already_indexed":False}
    finally:
        try: tmp.unlink(missing_ok=True)
        except OSError: pass


def _stored_path(path:str|Path,relative:str)->Path:
    rel=PurePosixPath(relative)
    if rel.is_absolute() or not rel.parts or ".." in rel.parts or "\\" in relative:
        raise ValueError("资料存储引用不是安全的相对路径")
    base=_db_path(path).parent; candidate=base.joinpath(*rel.parts); _safe_path(candidate,"已存储资料")
    if candidate!=base and base not in candidate.parents: raise ValueError("资料存储引用越过数据库目录")
    return candidate


def _tokens(value:str)->list[str]:
    normalized=re.sub(r"[，,;；|]+"," ",value.strip())
    return list(dict.fromkeys(re.findall(r"[A-Za-z0-9][A-Za-z0-9._/-]*|[\u3400-\u9fff]+",normalized)))


def _looks_like_model(value:str)->bool:
    value=value.strip()
    return bool(value and len(value)<=80 and re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9.\-]*",value)
        and re.search(r"[A-Za-z]",value) and re.search(r"\d",value))


def _exact_model(text:str,model:str)->bool:
    return re.search(r"(?<![A-Z0-9.\-])"+re.escape(model.upper())+r"(?![A-Z0-9.\-])",text.upper()) is not None


def _snippet(text:str,tokens:list[str],limit:int=260)->str:
    if len(text)<=limit: return text
    lower=text.casefold(); positions=[lower.find(t.casefold()) for t in tokens if t and lower.find(t.casefold())>=0]
    start=max(0,min(positions)-limit//4) if positions else 0; end=min(len(text),start+limit)
    return ("…" if start else "")+text[start:end]+("…" if end<len(text) else "")


def _obj(value:str)->dict[str,Any]:
    try:
        result=json.loads(value)
        return result if isinstance(result,dict) else {}
    except (ValueError,TypeError): return {}


def search_models(path:str|Path=DEFAULT_DB,*,query:str="",limit:int=100)->list[dict[str,Any]]:
    """Search registered manufacturer model identities; free-text SKU models stay separate."""
    if type(limit) is not int or not 1<=limit<=200: raise ValueError("limit 必须在 1 到 200 之间")
    if not isinstance(query,str) or len(query)>200: raise ValueError("型号查询条件无效")
    tokens=_tokens(query)
    with open_knowledge_readonly(path) as conn:
        validate_ingestion(conn)
        sql="""SELECT m.id,m.manufacturer_name,m.model_code,m.category,m.series,m.created_utc,
          (SELECT COUNT(*) FROM catalog_sku_model_links l WHERE l.model_id=m.id),
          (SELECT e.state FROM catalog_sku_model_events e JOIN catalog_sku_model_links l ON l.id=e.link_id
           WHERE l.model_id=m.id ORDER BY e.event_id DESC LIMIT 1)
          FROM catalog_models m WHERE 1=1"""
        args=[]
        for token in tokens:
            sql+=" AND INSTR(UPPER(m.manufacturer_name||' '||m.model_code||' '||m.category||' '||m.series),UPPER(?))>0"
            args.append(token)
        sql+=" ORDER BY m.manufacturer_name,m.model_code LIMIT ?"; args.append(limit)
        rows=conn.execute(sql,args).fetchall()
    return [{"type":"model","model_id":r[0],"manufacturer":r[1],"model_code":r[2],"category":r[3],
        "series":r[4],"created_utc":r[5],"sku_mapping_count":r[6],"latest_mapping_state":r[7],
        "identity_note":"制造商型号档案；只有显式确认的 SKU 映射才建立关联"} for r in rows]


def _document_results(conn:sqlite3.Connection,tokens:list[str],model:str,limit:int)->list[dict[str,Any]]:
    sql="""SELECT d.id,d.source_id,d.segment_type,d.ordinal,d.pdf_page,d.printed_page,d.sheet_name,d.cell_range,
        d.line_start,d.line_end,d.heading_path,d.text,d.region_json,s.original_filename,s.source_kind,
        s.authority_level,s.sha256,s.storage_ref FROM knowledge_document_segments d
        JOIN knowledge_sources s ON s.id=d.source_id WHERE d.parser_version=?"""
    args=[PARSER_VERSION]
    for token in tokens: sql+=" AND INSTR(UPPER(d.text),UPPER(?))>0"; args.append(token)
    if model: sql+=" AND INSTR(UPPER(d.text),UPPER(?))>0"; args.append(model)
    sql+=" ORDER BY d.source_id,d.pdf_page,d.sheet_name,d.ordinal LIMIT ?"; args.append(min(max(limit*30,300),5000))
    rows=conn.execute(sql,args).fetchall(); candidates=[]; precise=set()
    for r in rows:
        text_value=r[11]
        if model and not _exact_model(text_value,model): continue
        location={"pdf_page":r[4],"printed_page":r[5] or None,"sheet_name":r[6] or None,
            "cell_range":r[7] or None,"line_start":r[8],"line_end":r[9],
            "heading_path":r[10] or None,"region":_obj(r[12])}
        group=(r[1],r[4],r[6] or None,r[10] or None)
        if r[2] in _PRECISE_TYPES: precise.add(group)
        score=20+10*len(tokens)+(25 if model else 0)+(15 if r[2] in _PRECISE_TYPES else 0)
        if tokens and all(t.casefold() in text_value.casefold() for t in tokens): score+=10
        candidates.append({"type":"document","source_id":r[1],"filename":r[13],"source_kind":r[14],
            "authority_level":r[15],"source_sha256":r[16],"storage_ref":r[17] or None,
            "segment_id":r[0],"segment_type":r[2],"text":text_value,
            "snippet":_snippet(text_value,tokens+([model] if model else [])),"location":location,"score":score,
            "verification_note":"这是原始资料的文字检索结果，不代表参数已通过独立验证"})
    results=[]
    for item in candidates:
        if item["segment_type"] in {"pdf_page","markdown_section","docx_section","text_block"}:
            loc=item["location"]
            if (item["source_id"],loc["pdf_page"],loc["sheet_name"],loc["heading_path"]) in precise: continue
        results.append(item)
    results.sort(key=lambda x:(-x["score"],x["filename"],x["location"]["pdf_page"] or 0,x["location"]["line_start"] or 0))
    return results[:limit]


def _claim_results(conn:sqlite3.Connection,tokens:list[str],model:str,limit:int)->list[dict[str,Any]]:
    identity="""COALESCE(m.model_code,'')||' '||COALESCE(s.model,'')||' '||
      COALESCE((SELECT cm.model_code FROM catalog_sku_model_links l JOIN catalog_models cm ON cm.id=l.model_id
       WHERE l.sku_id=c.subject_id AND c.subject_type='sku'
       AND (SELECT ev.state FROM catalog_sku_model_events ev WHERE ev.link_id=l.id ORDER BY ev.event_id DESC LIMIT 1)='confirmed'
       ORDER BY l.created_utc LIMIT 1),'')"""
    combined=f"COALESCE(c.attribute,'')||' '||COALESCE(c.literal_value,'')||' '||COALESCE(c.unit,'')||' '||{identity}"
    sql=f"""SELECT c.id,c.subject_type,c.subject_id,c.attribute,c.literal_value,c.unit,c.claim_kind,c.normalized_value_json,
      c.created_by,c.created_utc,(SELECT e.state FROM knowledge_claim_events e WHERE e.claim_id=c.id
      ORDER BY e.event_id DESC LIMIT 1),COALESCE(m.model_code,''),COALESCE(s.model,''),COALESCE(s.code,''),
      COALESCE(p.name,''),COALESCE(p.category,''),COALESCE((SELECT cm.model_code FROM catalog_sku_model_links l
      JOIN catalog_models cm ON cm.id=l.model_id WHERE l.sku_id=c.subject_id AND c.subject_type='sku'
      AND (SELECT ev.state FROM catalog_sku_model_events ev WHERE ev.link_id=l.id ORDER BY ev.event_id DESC LIMIT 1)='confirmed'
      ORDER BY l.created_utc LIMIT 1),'')
      FROM knowledge_claims c LEFT JOIN catalog_models m ON c.subject_type='model' AND m.id=c.subject_id
      LEFT JOIN catalog_skus s ON c.subject_type='sku' AND s.id=c.subject_id
      LEFT JOIN catalog_products p ON s.product_id=p.id WHERE 1=1"""
    args=[]
    for token in tokens: sql+=f" AND INSTR(UPPER({combined}),UPPER(?))>0"; args.append(token)
    if model: sql+=f" AND INSTR(UPPER({identity}),UPPER(?))>0"; args.append(model)
    sql+=" ORDER BY c.created_utc DESC,c.id LIMIT ?"; args.append(min(max(limit*20,200),2000))
    rows=conn.execute(sql,args).fetchall(); results=[]
    for r in rows:
        identity_text=" ".join((r[11],r[12],r[16]))
        if model and not _exact_model(identity_text,model): continue
        binding=("confirmed_sku_model_mapping" if r[16] and _exact_model(r[16],model or r[16]) else
                 "sku_free_text_unverified" if r[12] else "no_model_text") if r[1]=="sku" else (
                 "manufacturer_model" if r[1]=="model" else "not_applicable")
        evrows=conn.execute("""SELECT e.id,e.source_id,src.original_filename,src.authority_level,e.evidence_role,
            e.pdf_page,e.printed_page,e.region_json,e.literal_excerpt,e.extraction_method,src.sha256
            FROM knowledge_claim_evidence e JOIN knowledge_sources src ON src.id=e.source_id
            WHERE e.claim_id=? ORDER BY e.created_utc,e.id""",(r[0],)).fetchall()
        evidence=[{"evidence_id":e[0],"source_id":e[1],"filename":e[2],"authority_level":e[3],"role":e[4],
            "pdf_page":e[5],"printed_page":e[6] or None,"region":_obj(e[7]),"literal_excerpt":e[8],
            "extraction_method":e[9],"source_sha256":e[10]} for e in evrows]
        results.append({"type":"claim","claim_id":r[0],"subject_type":r[1],"subject_id":r[2],
            "attribute":r[3],"literal_value":r[4],"unit":r[5],"claim_kind":r[6],
            "normalized_value_json":r[7],"created_by":r[8],"created_utc":r[9],"state":r[10] or "unknown",
            "model_code":r[11] or r[16] or r[12] or None,"sku_code":r[13] or None,
            "product_name":r[14] or None,"category":r[15] or None,"model_binding_status":binding,
            "evidence":evidence,"score":45+10*len(tokens)+(25 if model else 0),
            "verification_note":"主张状态与证据必须一起查看；accepted 仅表示通过数据库门槛"})
    return results[:limit]


def search_knowledge(path:str|Path=DEFAULT_DB,*,query:str="",model_code:str="",attribute:str="",limit:int=50)->dict[str,Any]:
    """Search source text, manufacturer models, SKU model text and existing claims."""
    if type(limit) is not int or not 1<=limit<=200: raise ValueError("limit 必须在 1 到 200 之间")
    for value,label,maxlen in ((query,"查询",300),(model_code,"型号",100),(attribute,"参数",150)):
        if not isinstance(value,str) or len(value)>maxlen: raise ValueError(f"{label}条件无效")
    model=model_code.strip()
    if not model and not attribute and _looks_like_model(query): model=query.strip()
    tokens=_tokens(" ".join(x for x in (query,attribute) if x))
    if not tokens and not model: raise ValueError("请提供关键词、型号或参数名称")
    with open_knowledge_readonly(path) as conn:
        validate_ingestion(conn)
        results=_document_results(conn,tokens,model,limit)+_claim_results(conn,tokens,model,limit)
        if model:
            rows=conn.execute("""SELECT m.id,m.manufacturer_name,m.model_code,m.category,m.series,
                (SELECT COUNT(*) FROM catalog_sku_model_links l WHERE l.model_id=m.id),
                (SELECT e.state FROM catalog_sku_model_events e JOIN catalog_sku_model_links l ON l.id=e.link_id
                 WHERE l.model_id=m.id ORDER BY e.event_id DESC LIMIT 1)
                FROM catalog_models m WHERE m.model_code=? COLLATE NOCASE ORDER BY m.manufacturer_name LIMIT ?""",
                (model,limit)).fetchall()
            results += [{"type":"model","model_id":r[0],"manufacturer":r[1],"model_code":r[2],"category":r[3],
                "series":r[4],"sku_mapping_count":r[5],"latest_mapping_state":r[6],"score":55,
                "identity_note":"制造商型号档案；不会自动绑定同名店铺 SKU"} for r in rows]
            rows=conn.execute("""SELECT s.id,s.code,s.model,p.name,p.category,st.name,
                COALESCE((SELECT cm.model_code FROM catalog_sku_model_links l JOIN catalog_models cm ON cm.id=l.model_id
                 WHERE l.sku_id=s.id AND (SELECT ev.state FROM catalog_sku_model_events ev WHERE ev.link_id=l.id
                 ORDER BY ev.event_id DESC LIMIT 1)='confirmed' ORDER BY l.created_utc LIMIT 1),'')
                FROM catalog_skus s JOIN catalog_products p ON p.id=s.product_id
                JOIN catalog_stores st ON st.id=s.store_id
                WHERE INSTR(UPPER(s.model),UPPER(?))>0 OR EXISTS(
                 SELECT 1 FROM catalog_sku_model_links l JOIN catalog_models cm ON cm.id=l.model_id
                 WHERE l.sku_id=s.id AND cm.model_code=? COLLATE NOCASE
                 AND (SELECT ev.state FROM catalog_sku_model_events ev WHERE ev.link_id=l.id
                      ORDER BY ev.event_id DESC LIMIT 1)='confirmed')
                ORDER BY st.name,p.name,s.code LIMIT ?""",(model,model,limit*5)).fetchall()
            for r in rows:
                raw=bool(r[2]) and _exact_model(r[2],model)
                confirmed=bool(r[6]) and r[6].casefold()==model.casefold()
                if raw or confirmed:
                    results.append({"type":"sku_model_text","sku_id":r[0],"sku_code":r[1],"model_text":r[2],
                        "product_name":r[3],"category":r[4],"store_name":r[5],
                        "model_binding_status":"confirmed_sku_model_mapping" if confirmed else "sku_free_text_unverified",
                        "score":35 if confirmed else 25,
                        "identity_note":"SKU 自由文本不是制造商型号身份；只有显式确认的映射才算已关联"})
    results.sort(key=lambda x:(-x.get("score",0),x.get("type","")))
    return {"query":query,"model_code":model or None,"attribute":attribute or None,
        "result_count":min(len(results),limit),"results":results[:limit],
        "search_note":"原始资料检索不自动生成参数主张；来源可信等级不是系统认证；型号文本不自动建立 SKU 关联"}


def knowledge_document_summary(path:str|Path=DEFAULT_DB)->dict[str,Any]:
    with open_knowledge_readonly(path) as conn:
        validate_ingestion(conn)
        counts={"sources":conn.execute("SELECT COUNT(*) FROM knowledge_sources").fetchone()[0],
            "indexed_sources":conn.execute("SELECT COUNT(DISTINCT source_id) FROM knowledge_ingest_runs WHERE status='indexed'").fetchone()[0],
            "segments":conn.execute("SELECT COUNT(*) FROM knowledge_document_segments").fetchone()[0],
            "indexed_runs":conn.execute("SELECT COUNT(*) FROM knowledge_ingest_runs WHERE status='indexed'").fetchone()[0],
            "failed_runs":conn.execute("SELECT COUNT(*) FROM knowledge_ingest_runs WHERE status='failed'").fetchone()[0]}
        formats={}
        for (metadata,) in conn.execute("SELECT metadata_json FROM knowledge_ingest_runs WHERE status='indexed'"):
            fmt=_obj(metadata).get("file_format","unknown"); formats[fmt]=formats.get(fmt,0)+1
    return {"schema_version":INGEST_SCHEMA_VERSION,"parser_version":PARSER_VERSION,"counts":counts,"indexed_formats":formats}


def audit_document_store(path:str|Path=DEFAULT_DB)->dict[str,Any]:
    """Read-only check of source bytes, segment digests, source locators and database integrity."""
    findings=[]; db=_db_path(path)
    with open_knowledge_readonly(db) as conn:
        validate_ingestion(conn)
        if conn.execute("PRAGMA quick_check").fetchone()[0]!="ok": findings.append("SQLite 物理完整性检查失败")
        if conn.execute("PRAGMA foreign_key_check").fetchall(): findings.append("文档检索库存在外键异常")
        latest=conn.execute("""SELECT r.source_id,r.parser_version,r.segment_count,r.storage_relpath,r.run_id,
            s.sha256,s.original_filename FROM knowledge_ingest_runs r JOIN knowledge_sources s ON s.id=r.source_id
            JOIN (SELECT source_id,MAX(event_id) last_id FROM knowledge_ingest_runs
                  WHERE status='indexed' GROUP BY source_id) x ON x.last_id=r.event_id""").fetchall()
        for source_id,parser_version,expected_count,relpath,run_id,expected_sha,filename in latest:
            try:
                stored=_stored_path(db,relpath)
                if not stored.is_file(): findings.append(f"{filename}: 已索引的原始文件缺失"); continue
                actual_sha,_=_sha256_file(stored)
                if actual_sha!=expected_sha: findings.append(f"{filename}: 原始文件 SHA-256 与登记值不一致")
                actual_count=conn.execute("SELECT COUNT(*) FROM knowledge_document_segments WHERE source_id=? AND parser_version=?",
                    (source_id,parser_version)).fetchone()[0]
                if actual_count!=expected_count: findings.append(f"{filename}: 索引段数与导入记录不一致")
                for segment_id,text_value,digest,region in conn.execute(
                    "SELECT id,text,text_sha256,region_json FROM knowledge_document_segments WHERE source_id=? AND parser_version=?",
                    (source_id,parser_version)):
                    if hashlib.sha256(text_value.encode("utf-8")).hexdigest()!=digest:
                        findings.append(f"{filename}: 文本段 {segment_id[:10]} 指纹不一致")
                    try:
                        if not isinstance(json.loads(region),dict): raise ValueError("region")
                    except (ValueError,TypeError): findings.append(f"{filename}: 文本段 {segment_id[:10]} 定位数据无效")
            except (OSError,ValueError) as exc: findings.append(f"{filename}: 资料路径检查失败（{exc}）")
    return {"ok":not findings,"findings":findings[:100],"additional_findings":max(0,len(findings)-100),
        "summary":knowledge_document_summary(db)}
