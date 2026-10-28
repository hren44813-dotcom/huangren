from __future__ import annotations
import csv,hashlib,io,json,os,sqlite3,subprocess,sys,tempfile,unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]; sys.path.insert(0,str(ROOT/"src"))
import fitz
from docx import Document
from openpyxl import Workbook
from ecom_images.catalog_v12 import COLUMNS,import_catalog_csv,list_skus
from ecom_images.database_v9 import open_database,save_job
from ecom_images.production_ledger_v6 import new_record
from ecom_images.knowledge_schema_v13 import initialize_knowledge
from ecom_images.knowledge_v13 import add_model
from ecom_images.knowledge_claims_v13 import create_claim,add_claim_evidence,add_claim_review
from ecom_images.knowledge_ingest_v13 import (
    audit_document_store,import_document,initialize_ingestion,knowledge_document_summary,search_knowledge,
)

def catalog_csv():
    row={"店铺编码":"SHOP-1","店铺名称":"工业五金店","平台":"天猫","商品编码":"P-1","商品名称":"测试转轴",
    "商品类别":"阻尼转轴","SKU编码":"SKU-1","产品型号":"GL1-336","材质":"","表面处理":"","孔距(mm)":"","总长(mm)":"","备注":""}
    stream=io.StringIO(newline=""); writer=csv.DictWriter(stream,fieldnames=COLUMNS); writer.writeheader(); writer.writerow(row)
    return ("\ufeff"+stream.getvalue()).encode("utf-8")

class KnowledgeIngestV13Tests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name); self.db=self.root/"knowledge.sqlite3"
        save_job(self.db,new_record([("SKU.jpg",b"synthetic",{})],fingerprint="a"*64,platform="天猫",template="蓝白"))
        import_catalog_csv(self.db,catalog_csv(),dry_run=False)
        initialize_knowledge(self.db); initialize_ingestion(self.db)
        self.source_dir=self.root/"inputs"; self.source_dir.mkdir()
    def write_pdf(self,name="catalog.pdf",pages=None):
        path=self.source_dir/name; doc=fitz.open()
        for lines in pages or [["型号 GL1-336","扭矩 6 N.m","材质 1215"]]:
            page=doc.new_page(width=595,height=842)
            for i,line in enumerate(lines): page.insert_text((50,70+i*28),line)
        doc.save(path); doc.close(); return path
    def test_pdf_page_number_and_line_bbox_are_searchable(self):
        path=self.write_pdf(pages=[["GL1-336-A XGL1-336"],["型号 GL1-336","扭矩 6 N.m"]])
        result=import_document(self.db,file_path=path,source_kind="manufacturer_document")
        found=search_knowledge(self.db,model_code="GL1-336")
        docs=[x for x in found["results"] if x["type"]=="document"]
        self.assertTrue(docs); self.assertTrue(all(x["location"]["pdf_page"]==2 for x in docs))
        line=next(x for x in docs if x["segment_type"]=="pdf_line")
        region=line["location"]["region"]
        self.assertEqual(region["coordinate_system"],"pdf_points_top_left")
        self.assertEqual(len(region["bbox"]),4); self.assertEqual(len(region["normalized_bbox"]),4)
        self.assertEqual(result["sha256"],hashlib.sha256(path.read_bytes()).hexdigest())
    def test_model_plus_parameter_search_falls_back_to_page_context(self):
        path=self.write_pdf(pages=[["Model GL1-336","Torque 6 N.m"]])
        import_document(self.db,file_path=path,source_kind="manufacturer_document")
        result=search_knowledge(self.db,model_code="GL1-336",attribute="Torque")
        docs=[x for x in result["results"] if x["type"]=="document"]
        self.assertTrue(docs); self.assertTrue(any(x["location"]["pdf_page"]==1 for x in docs))
    def test_suffix_and_longer_code_are_not_exact_model_matches(self):
        path=self.write_pdf(pages=[["GL1-336-A","XGL1-336"]])
        import_document(self.db,file_path=path,source_kind="manufacturer_document")
        result=search_knowledge(self.db,model_code="GL1-336")
        self.assertFalse([x for x in result["results"] if x["type"]=="document"])
    def test_markdown_keeps_heading_and_line_locations(self):
        path=self.source_dir/"specs.md"
        path.write_text("# GL1 系列\n\n## 扭矩参数\n型号：GL1-336\n额定扭矩：6 N.m\n",encoding="utf-8")
        import_document(self.db,file_path=path,source_kind="manufacturer_document")
        result=search_knowledge(self.db,model_code="GL1-336",attribute="扭矩")
        section=next(x for x in result["results"] if x["type"]=="document" and x["segment_type"]=="markdown_section")
        self.assertIn("扭矩参数",section["location"]["heading_path"])
        self.assertEqual(section["location"]["line_start"],3)
        self.assertEqual(section["location"]["line_end"],5)
    def test_docx_keeps_paragraph_table_and_heading_locations(self):
        path=self.source_dir/"specs.docx"; doc=Document()
        doc.add_heading("GL1-336 规格",level=1); doc.add_paragraph("额定扭矩以原厂资料为准。")
        table=doc.add_table(rows=2,cols=2); table.cell(0,0).text="型号"; table.cell(0,1).text="扭矩"
        table.cell(1,0).text="GL1-336"; table.cell(1,1).text="6 N.m"; doc.save(path)
        imported=import_document(self.db,file_path=path,source_kind="manufacturer_document")
        result=search_knowledge(self.db,model_code="GL1-336",attribute="扭矩")
        rows=[x for x in result["results"] if x["type"]=="document" and x["segment_type"]=="docx_table_row"]
        self.assertTrue(rows); self.assertEqual(rows[0]["location"]["sheet_name"],"Table 1")
        self.assertIn("R2C1:R2C2",rows[0]["location"]["cell_range"])
        self.assertIsNone(rows[0]["location"]["pdf_page"])
        self.assertTrue(any("DOCX" in warning for warning in imported["warnings"]))
    def test_xlsx_header_values_and_cell_range_are_searchable(self):
        path=self.source_dir/"specs.xlsx"; wb=Workbook(); ws=wb.active; ws.title="GL1参数"
        ws.append(["型号","扭矩"]); ws.append(["GL1-336","6 N.m"]); wb.save(path)
        import_document(self.db,file_path=path,source_kind="manufacturer_document")
        result=search_knowledge(self.db,model_code="GL1-336",attribute="扭矩")
        row=next(x for x in result["results"] if x["type"]=="document" and x["segment_type"]=="excel_row")
        self.assertEqual(row["location"]["sheet_name"],"GL1参数")
        self.assertEqual(row["location"]["cell_range"],"A2:B2")
        self.assertEqual(row["location"]["region"]["cells"][1]["cell"],"B2")
    def test_txt_line_numbers_are_preserved(self):
        path=self.source_dir/"notes.txt"; path.write_text("说明\n\n型号 GL1-336\n扭矩 6 N.m\n",encoding="utf-8")
        import_document(self.db,file_path=path)
        result=search_knowledge(self.db,model_code="GL1-336")
        item=next(x for x in result["results"] if x["type"]=="document" and x["segment_type"]=="text_line")
        self.assertEqual(item["location"]["line_start"],3)
    def test_reimport_is_idempotent_and_does_not_create_claims(self):
        path=self.write_pdf()
        first=import_document(self.db,file_path=path,source_kind="manufacturer_document")
        before=knowledge_document_summary(self.db)
        second=import_document(self.db,file_path=path,source_kind="manufacturer_document")
        after=knowledge_document_summary(self.db)
        self.assertEqual(first["run_id"],second["run_id"]); self.assertTrue(second["already_indexed"])
        self.assertEqual(before["counts"]["segments"],after["counts"]["segments"])
        with open_database(self.db) as conn: self.assertEqual(conn.execute("SELECT COUNT(*) FROM knowledge_claims").fetchone()[0],0)
    def test_sku_free_text_is_not_treated_as_confirmed_mapping(self):
        result=search_knowledge(self.db,model_code="GL1-336")
        sku=next(x for x in result["results"] if x["type"]=="sku_model_text")
        self.assertEqual(sku["model_binding_status"],"sku_free_text_unverified")
        self.assertEqual(sku["model_text"],"GL1-336")
    def test_claim_search_returns_original_evidence_locator(self):
        path=self.write_pdf(); imported=import_document(self.db,file_path=path,source_kind="manufacturer_document")
        model_id=add_model(self.db,model_code="GL1-336",category="阻尼转轴",manufacturer_name="测试制造商")
        claim=create_claim(self.db,subject_type="model",subject_id=model_id,attribute="扭矩",literal_value="6",unit="N.m",
            claim_kind="catalog_transcription",created_by="测试")
        add_claim_evidence(self.db,claim_id=claim,source_id=imported["source_id"],evidence_role="supports",
            extraction_method="pdf_text",created_by="测试",pdf_page=1,region={"bbox":[50,60,160,80]},literal_excerpt="扭矩 6 N.m")
        add_claim_review(self.db,claim_id=claim,verdict="supports",reviewer_role="second_reviewer",
            reviewer_name="复核人",method="对照原文",notes="测试数据")
        result=search_knowledge(self.db,model_code="GL1-336",attribute="扭矩")
        item=next(x for x in result["results"] if x["type"]=="claim")
        self.assertEqual(item["state"],"proposed")
        self.assertEqual(item["evidence"][0]["pdf_page"],1)
        self.assertEqual(item["evidence"][0]["region"]["bbox"],[50,60,160,80])
    def test_audit_passes_for_unchanged_import_and_detects_tampering(self):
        path=self.write_pdf(); result=import_document(self.db,file_path=path)
        self.assertTrue(audit_document_store(self.db)["ok"])
        stored=self.root/result["storage_relpath"]; stored.write_bytes(stored.read_bytes()+b"tamper")
        audit=audit_document_store(self.db)
        self.assertFalse(audit["ok"]); self.assertTrue(any("SHA-256" in x for x in audit["findings"]))
    def test_unsupported_legacy_doc_is_rejected(self):
        path=self.source_dir/"legacy.doc"; path.write_bytes(b"not a doc")
        with self.assertRaisesRegex(ValueError,"另存为 .docx"): import_document(self.db,file_path=path)
        with open_database(self.db) as conn: self.assertEqual(conn.execute("SELECT COUNT(*) FROM knowledge_sources").fetchone()[0],0)
    def test_invalid_pdf_is_rejected_before_source_registration(self):
        path=self.source_dir/"bad.pdf"; path.write_bytes(b"not a PDF")
        with self.assertRaisesRegex(ValueError,"无法打开 PDF"): import_document(self.db,file_path=path)
        with open_database(self.db) as conn: self.assertEqual(conn.execute("SELECT COUNT(*) FROM knowledge_sources").fetchone()[0],0)
    def test_source_size_limit_is_enforced(self):
        path=self.source_dir/"notes.md"; path.write_text("A"*100,encoding="utf-8")
        with self.assertRaisesRegex(ValueError,"大小限制"): import_document(self.db,file_path=path,max_bytes=10)
    def test_symlink_source_is_rejected(self):
        target=self.source_dir/"real.md"; target.write_text("型号 GL1-336",encoding="utf-8")
        link=self.source_dir/"link.md"
        try: link.symlink_to(target)
        except (OSError,NotImplementedError): self.skipTest("当前环境不允许创建符号链接")
        with self.assertRaisesRegex(ValueError,"符号链接"): import_document(self.db,file_path=link)
    def test_cli_init_import_search_and_audit(self):
        path=self.write_pdf(); cli=[sys.executable,str(ROOT/"scripts/knowledge_ingest_v13.py"),"--database",str(self.db)]
        init=subprocess.run(cli+["init"],cwd=ROOT,text=True,capture_output=True,timeout=30)
        self.assertEqual(init.returncode,0,init.stdout+init.stderr)
        imp=subprocess.run(cli+["import-file","--file",str(path),"--kind","manufacturer_document"],cwd=ROOT,text=True,capture_output=True,timeout=30)
        self.assertEqual(imp.returncode,0,imp.stdout+imp.stderr); self.assertIn('"source_id"',imp.stdout)
        search=subprocess.run(cli+["search","--model","GL1-336"],cwd=ROOT,text=True,capture_output=True,timeout=30)
        self.assertEqual(search.returncode,0,search.stdout+search.stderr); self.assertIn("pdf_page",search.stdout)
        audit=subprocess.run(cli+["audit"],cwd=ROOT,text=True,capture_output=True,timeout=30)
        self.assertEqual(audit.returncode,0,audit.stdout+audit.stderr)
    def test_read_only_search_does_not_create_missing_database(self):
        missing=self.root/"missing.sqlite3"
        with self.assertRaises(ValueError): search_knowledge(missing,query="扭矩")
        self.assertFalse(missing.exists())

if __name__=="__main__": unittest.main(verbosity=2)
